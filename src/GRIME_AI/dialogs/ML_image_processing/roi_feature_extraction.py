#!/usr/bin/env python3
# -*- coding: utf-8 -*-

# Author: John Edward Stranzl, Jr.
# Affiliation(s): University of Nebraska-Lincoln, Blade Vision Systems, LLC
# Contact: jstranzl2@huskers.unl.edu, johnstranzl@gmail.com
# License: Apache License, Version 2.0, http://www.apache.org/licenses/LICENSE-2.0

# roi_feature_extraction.py
#
# ROI feature extraction shared by the ROI Analyzer tab and the Segment Images
# tab, so both produce identical columns. Every feature is computed inside the
# segmentation mask. No Qt widgets here.
#
#   RoiFeatureExtractor(settings).header()                  column names
#   RoiFeatureExtractor(settings).compute_row(img, mask)    one row, never raises
#   export_pairs(pairs, folder, settings, ...)              CSV (+ XLSX) for many pairs
#   segmentation_pairs(input_dir, output_dir, ...)          pairs after segmentation
#   load_settings() / save_settings(settings)               persisted options

import os
import re
import csv
import copy
import json
from datetime import datetime

import cv2
import numpy as np

from appcore.utils.fft_utils import fft_radial_profile, DEFAULT_RADIAL_RINGS
from appcore.Texture import GLCM_NEIGHBOR_ANGLES

SETTINGS_KEY = "roi_feature_extraction"
SETTINGS_FILE = "site_config.json"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
STATUS_OK = "OK"

DEFAULT_SETTINGS = {
    "clustering": {
        "method": "kmeans",            # kmeans | gmm | meanshift
        "kmeans_clusters": 3,
        "gmm_clusters": 3,
        "meanshift_auto": True,
        "meanshift_quantile": 0.2,
        "meanshift_bandwidth": 15.0,
    },
    "texture": {
        "glcm":    {"enabled": True,  "distance": 1},
        "gabor":   {"enabled": False, "frequencies": [0.1, 0.3]},
        "lbp":     {"enabled": False, "points": 8, "radius": 1},
        "wavelet": {"enabled": False, "wavelet": "db1", "levels": 2},
        "fourier": {"enabled": False, "rings": DEFAULT_RADIAL_RINGS},
    },
    "output": {
        "excel": True,
    },
    "sensor": {
        "enabled": False,          # correlate NWIS sensor data into the feature file
        "auto_detect": True,       # look in the sister "data" folder beside the images
        "sensor_file": "",         # used when auto_detect is off, or once chosen
        "tolerance_minutes": 0,    # 0: automatic (half the sensor sampling interval)
    },
}

# The image download creates <root>/images and <root>/data; the NWIS file lands in data.
SENSOR_DATA_FOLDER_NAME = "data"
# Downloads arrive as USGS text and are converted to CSV; only the CSV is used, so
# the two copies of one download are not offered as a choice.
SENSOR_FILE_EXTENSIONS = (".csv",)

GLCM_PROPERTIES = ("contrast", "homogeneity", "correlation")
WAVELET_BANDS = ("cH", "cV", "cD")


# ======================================================================================================================
# Settings
# ======================================================================================================================
def merged_settings(settings=None) -> dict:
    """DEFAULT_SETTINGS with any given values laid over it (missing keys keep defaults)."""
    def merge(base, over):
        out = copy.deepcopy(base)
        for k, v in (over or {}).items():
            out[k] = merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
        return out
    return merge(DEFAULT_SETTINGS, settings)


def _settings_path():
    from appcore.Save_Utils import Save_Utils
    return os.path.join(os.path.normpath(Save_Utils().get_settings_folder()), SETTINGS_FILE)


def load_settings() -> dict:
    try:
        with open(_settings_path(), "r") as f:
            return merged_settings(json.load(f).get(SETTINGS_KEY))
    except Exception:
        return merged_settings()


def save_settings(settings: dict):
    path = _settings_path()
    config = {}
    if os.path.exists(path):
        with open(path, "r") as f:
            config = json.load(f)
    config[SETTINGS_KEY] = merged_settings(settings)
    with open(path, "w") as f:
        json.dump(config, f, indent=4)


def clustering_kwargs(settings) -> dict:
    """Keyword arguments for ROI_Analyzer from the clustering settings."""
    c = merged_settings(settings)["clustering"]
    method = c["method"]
    kwargs = {"clusters": cluster_count(settings), "clustering_method": method}
    if method == "meanshift":
        if c["meanshift_auto"]:
            kwargs.update(bandwidth=None, quantile=c["meanshift_quantile"])
        else:
            kwargs.update(bandwidth=c["meanshift_bandwidth"], quantile=None)
    return kwargs


def cluster_count(settings) -> int:
    """Number of cluster column groups. Mean-shift finds its own count; the K-Means
    count sets how many of its clusters are reported (as before)."""
    c = merged_settings(settings)["clustering"]
    return int(c["gmm_clusters"] if c["method"] == "gmm" else c["kmeans_clusters"])


def clustering_params_text(settings) -> str:
    c = merged_settings(settings)["clustering"]
    if c["method"] == "meanshift":
        if c["meanshift_auto"]:
            return f"auto=True, quantile={c['meanshift_quantile']}"
        return f"auto=False, bandwidth={c['meanshift_bandwidth']}"
    return f"clusters={cluster_count(settings)}"


def sister_data_folder(images_folder) -> str:
    """The data folder beside the images folder, from the download layout."""
    return os.path.join(os.path.dirname(os.path.normpath(images_folder)), SENSOR_DATA_FOLDER_NAME)


def sensor_files_in(folder) -> list:
    """NWIS CSV files in folder, newest first. Repeated downloads of the same range
    leave more than one, and then the user is asked which to use."""
    if not folder or not os.path.isdir(folder):
        return []
    paths = [os.path.join(folder, f) for f in os.listdir(folder)
             if f.lower().endswith(SENSOR_FILE_EXTENSIONS)]
    return sorted(paths, key=lambda p: os.path.getmtime(p), reverse=True)


# ======================================================================================================================
# Capture date/time from filename
# ======================================================================================================================
_DATETIME_PATTERNS = [
    (r'(\d{4})-(\d{2})-(\d{2})T(\d{2})-(\d{2})-(\d{2})Z', True),
    (r'(\d{4})_(\d{2})_(\d{2})_(\d{2})(\d{2})(\d{2})',    True),
    (r'(\d{4})(\d{2})(\d{2})_(\d{2})(\d{2})(\d{2})',      True),
    (r'(\d{4})(\d{2})(\d{2})(\d{2})(\d{2})(\d{2})',       True),
    (r'(\d{4})_(\d{2})_(\d{2})',                          False),
    (r'(\d{4})(\d{2})(\d{2})',                            False),
]


def extract_datetime_from_path(filepath):
    """(date 'YYYY-MM-DD', time 'HH:MM:SS') from the filename, 'n/a' for parts not found."""
    stem = os.path.splitext(os.path.basename(filepath))[0]
    for pattern, has_time in _DATETIME_PATTERNS:
        m = re.search(pattern, stem)
        if m:
            g = m.groups()
            date = f"{g[0]}-{g[1]}-{g[2]}"
            return (date, f"{g[3]}:{g[4]}:{g[5]}") if has_time else (date, "n/a")
    return "n/a", "n/a"


# ======================================================================================================================
# Masked texture features
# ======================================================================================================================
def _crop_to_mask(gray, mask):
    """Crop gray and mask to the mask's bounding box. Returns (gray, mask) or (None, None)."""
    ys, xs = np.nonzero(mask)
    if ys.size == 0:
        return None, None
    y0, y1, x0, x1 = ys.min(), ys.max() + 1, xs.min(), xs.max() + 1
    return gray[y0:y1, x0:x1], mask[y0:y1, x0:x1]


def masked_glcm(gray, mask, distance):
    """GLCM contrast, homogeneity and correlation from pixel pairs inside the mask only,
    averaged over the four neighbor directions."""
    from skimage.feature import graycomatrix, graycoprops
    g, m = _crop_to_mask(gray, mask)
    if g is None:
        return {p: np.nan for p in GLCM_PROPERTIES}
    # Out-of-mask pixels get an extra gray level (256); its row and column are
    # dropped, so every remaining pair has both pixels inside the mask.
    img = g.astype(np.uint16)
    img[~m] = 256
    P = graycomatrix(img, [int(distance)], GLCM_NEIGHBOR_ANGLES, levels=257,
                     symmetric=True, normed=False)[:256, :256].astype(np.float64)
    totals = P.sum(axis=(0, 1), keepdims=True)
    valid = totals[0, 0] > 0
    if not valid.any():
        return {p: np.nan for p in GLCM_PROPERTIES}
    P = np.divide(P, totals, out=np.zeros_like(P), where=totals > 0)
    out = {}
    for prop in GLCM_PROPERTIES:
        values = graycoprops(P, prop)[valid]
        out[prop] = float(np.nanmean(values)) if values.size else np.nan
    return out


def masked_gabor(gray, mask, frequencies):
    """Mean and variance of the Gabor magnitude inside the mask for each frequency,
    averaged over the four neighbor directions."""
    from skimage.filters import gabor
    g, m = _crop_to_mask(gray, mask)
    out = {}
    for f in frequencies:
        if g is None:
            out[f] = (np.nan, np.nan)
            continue
        means, variances = [], []
        for theta in GLCM_NEIGHBOR_ANGLES:
            real, imag = gabor(g.astype(np.float64), frequency=float(f), theta=theta)
            mag = np.hypot(real, imag)[m]
            means.append(mag.mean())
            variances.append(mag.var())
        out[f] = (float(np.mean(means)), float(np.mean(variances)))
    return out


def masked_lbp(gray, mask, points, radius):
    """Normalized histogram of uniform LBP codes inside the mask (points + 2 bins)."""
    from skimage.feature import local_binary_pattern
    n_bins = int(points) + 2
    g, m = _crop_to_mask(gray, mask)
    if g is None:
        return [np.nan] * n_bins
    codes = local_binary_pattern(g, int(points), float(radius), method="uniform")[m]
    hist = np.bincount(codes.astype(np.int64), minlength=n_bins)[:n_bins].astype(np.float64)
    return list(hist / hist.sum()) if hist.sum() > 0 else [np.nan] * n_bins


def masked_wavelet(gray, mask, wavelet, levels):
    """Variance of the detail coefficients inside the mask, keyed (level, band).
    Level 1 is the finest."""
    import pywt
    levels = int(levels)
    g, m = _crop_to_mask(gray, mask)
    out = {(lvl, b): np.nan for lvl in range(1, levels + 1) for b in WAVELET_BANDS}
    if g is None:
        return out
    coeffs = pywt.wavedec2(g.astype(np.float64), wavelet, level=levels)
    m8 = m.astype(np.uint8)
    for k, bands in enumerate(coeffs[1:]):
        lvl = levels - k                       # coeffs[1] is the coarsest level
        h, w = bands[0].shape
        mk = cv2.resize(m8, (w, h), interpolation=cv2.INTER_NEAREST) > 0
        if not mk.any():
            continue
        for name, band in zip(WAVELET_BANDS, bands):
            out[(lvl, name)] = float(band[mk].var())
    return out


def masked_fourier(gray, mask, rings):
    """Radial FFT magnitude profile of the region. Pixels outside the mask are set to
    the region's mean so they add no detail of their own."""
    g, m = _crop_to_mask(gray, mask)
    if g is None or min(g.shape) < 2:
        return [np.nan] * int(rings)
    region = g.astype(np.float64)
    region[~m] = region[m].mean()
    return list(fft_radial_profile(region, num_bins=int(rings)))


# ======================================================================================================================
# Extractor
# ======================================================================================================================
BASE_COLUMNS = [
    "Image Path", "Mask Path", "Capture Date", "Capture Time",
    "Clustering Method", "Clustering Params",
    "ROI Intensity", "ROI Entropy", "ROI Texture", "Mean GLI", "Mean GCC",
    "GLCM Contrast", "GLCM Homogeneity", "GLCM Correlation",
    "ROI Pixel Count", "ROI Area", "Image Height", "Image Width",
    "Image Total Pixels", "ROI Area Percentage",
]


def _fmt(value, decimals):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    return f"{value:.{decimals}f}"


class RoiFeatureExtractor:

    def __init__(self, settings=None):
        self.settings = merged_settings(settings)
        self.n_clusters = cluster_count(self.settings)

    # ------------------------------------------------------------------------------------------------------------------
    def header(self) -> list:
        t = self.settings["texture"]
        cols = list(BASE_COLUMNS)
        cols += [f"Cluster {i} {c}" for i in range(1, self.n_clusters + 1) for c in "HSV"]
        cols += [f"Cluster {i} {c}" for i in range(1, self.n_clusters + 1) for c in "RGB"]
        if t["gabor"]["enabled"]:
            for f in t["gabor"]["frequencies"]:
                cols += [f"Gabor f={f} Mean", f"Gabor f={f} Var"]
        if t["lbp"]["enabled"]:
            cols += [f"LBP Bin {i}" for i in range(int(t["lbp"]["points"]) + 2)]
        if t["wavelet"]["enabled"]:
            cols += [f"Wavelet L{lvl} {b} Var"
                     for lvl in range(1, int(t["wavelet"]["levels"]) + 1) for b in WAVELET_BANDS]
        if t["fourier"]["enabled"]:
            cols += [f"Fourier Ring {i}" for i in range(1, int(t["fourier"]["rings"]) + 1)]
        cols.append("Status")
        return cols

    # ------------------------------------------------------------------------------------------------------------------
    def compute_row(self, image_path, mask_path) -> list:
        """One row for header(). Never raises: a failure gives a row with the paths,
        empty feature cells and the reason in Status."""
        capture_date, capture_time = extract_datetime_from_path(image_path)
        lead = [image_path, mask_path or "", capture_date, capture_time,
                self.settings["clustering"]["method"], clustering_params_text(self.settings)]
        width = len(self.header())
        try:
            if not mask_path or not os.path.exists(mask_path):
                raise ValueError("No mask for this image (target not found or mask not saved)")
            values = self._features(image_path, mask_path)
        except Exception as e:
            reason = str(e)
            return lead + [""] * (width - len(lead) - 1) + [f"Failed: {reason}"]
        return lead + values + [STATUS_OK]

    # ------------------------------------------------------------------------------------------------------------------
    def _features(self, image_path, mask_path) -> list:
        from appcore.ROI_Analyzer import ROI_Analyzer

        mask_check = cv2.imread(mask_path, cv2.IMREAD_GRAYSCALE)
        if mask_check is None:
            raise ValueError("Mask could not be read")
        if not np.any(mask_check > 1):
            raise ValueError("Mask is empty (target not found in this image)")

        a = ROI_Analyzer(image_path, mask_path, **clustering_kwargs(self.settings))
        a.run_analysis()

        gray = cv2.cvtColor(a.image, cv2.COLOR_BGR2GRAY)
        mask = a.mask_bin > 0
        t = self.settings["texture"]

        glcm = (masked_glcm(gray, mask, t["glcm"]["distance"]) if t["glcm"]["enabled"]
                else {p: None for p in GLCM_PROPERTIES})

        row = [
            _fmt(a.roi_intensity, 2), _fmt(a.roi_entropy, 4), _fmt(a.roi_texture, 4),
            _fmt(a.mean_gli, 6), _fmt(a.mean_gcc, 6),
            _fmt(glcm["contrast"], 4), _fmt(glcm["homogeneity"], 4), _fmt(glcm["correlation"], 4),
            _fmt(a.ROI_total_pixels, 2), _fmt(a.ROI_total_area, 2),
            _fmt(a.image_height, 2), _fmt(a.image_width, 2),
            _fmt(a.image_total_pixels, 2), _fmt(a.ROI_percentage, 2),
        ]

        def fixed(seq):     # always n_clusters groups, so columns line up for every image
            seq = list(seq)[:self.n_clusters]
            return seq + [(None, None, None)] * (self.n_clusters - len(seq))
        for triple in fixed(a.dominant_hsv_list) + fixed(a.dominant_rgb_list):
            row += [_fmt(v, 4) for v in triple]

        if t["gabor"]["enabled"]:
            for f, (mean, var) in masked_gabor(gray, mask, t["gabor"]["frequencies"]).items():
                row += [_fmt(mean, 4), _fmt(var, 4)]
        if t["lbp"]["enabled"]:
            row += [_fmt(v, 6) for v in masked_lbp(gray, mask, t["lbp"]["points"], t["lbp"]["radius"])]
        if t["wavelet"]["enabled"]:
            wv = masked_wavelet(gray, mask, t["wavelet"]["wavelet"], t["wavelet"]["levels"])
            row += [_fmt(wv[(lvl, b)], 4)
                    for lvl in range(1, int(t["wavelet"]["levels"]) + 1) for b in WAVELET_BANDS]
        if t["fourier"]["enabled"]:
            row += [_fmt(v, 6) for v in masked_fourier(gray, mask, t["fourier"]["rings"])]
        return row


# ======================================================================================================================
# Output
# ======================================================================================================================
def run_prefix() -> str:
    """Date/time filename prefix used for every run, e.g. 20260922_101500."""
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def write_xlsx(xlsx_path, header, rows):
    """Workbook with the same columns as the CSV. Image and mask paths become
    hyperlinks showing the file name; numeric text becomes numbers."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "ROI Metrics"
    for c, title in enumerate(header, start=1):
        ws.cell(row=1, column=c, value=title)

    link_cols = {header.index(n) for n in ("Image Path", "Mask Path") if n in header}
    link_font = Font(color="0000FF", underline="single")
    for r, data in enumerate(rows, start=2):
        for c, value in enumerate(data):
            if c in link_cols and value:
                cell = ws.cell(row=r, column=c + 1, value=os.path.basename(str(value)))
                cell.hyperlink = str(value)
                cell.font = link_font
                continue
            try:
                value = float(value) if isinstance(value, str) and value not in ("", "n/a") else value
            except ValueError:
                pass
            ws.cell(row=r, column=c + 1, value=value)
    wb.save(xlsx_path)


def export_pairs(pairs, output_folder, settings=None, prefix=None,
                 progress_callback=None, cancel_callback=None):
    """
    Compute features for (image_path, mask_path) pairs and write
    <prefix>_roi_metrics.csv (row by row, so a stopped run keeps its rows) and,
    if enabled, <prefix>_roi_metrics.xlsx in output_folder.

    progress_callback(done, total) and cancel_callback() -> bool are optional.
    Returns dict: csv_path, xlsx_path (None if not written), rows_ok, rows_failed, cancelled.
    """
    settings = merged_settings(settings)
    extractor = RoiFeatureExtractor(settings)
    header = extractor.header()
    prefix = prefix or run_prefix()
    os.makedirs(output_folder, exist_ok=True)
    csv_path = os.path.join(output_folder, f"{prefix}_roi_metrics.csv")

    rows, n_ok, n_failed, cancelled = [], 0, 0, False
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for i, (image_path, mask_path) in enumerate(pairs):
            if cancel_callback is not None and cancel_callback():
                cancelled = True
                break
            row = extractor.compute_row(image_path, mask_path)
            writer.writerow(row)
            f.flush()
            rows.append(row)
            if row[-1] == STATUS_OK:
                n_ok += 1
            else:
                n_failed += 1
            if progress_callback is not None:
                progress_callback(i + 1, len(pairs))

    xlsx_path = None
    if settings["output"]["excel"] and rows:
        xlsx_path = os.path.join(output_folder, f"{prefix}_roi_metrics.xlsx")
        try:
            write_xlsx(xlsx_path, header, rows)
        except Exception as e:
            print(f"[roi_feature_extraction] Could not write {xlsx_path}: {e}")
            xlsx_path = None

    return {"csv_path": csv_path, "xlsx_path": xlsx_path, "rows_ok": n_ok,
            "rows_failed": n_failed, "cancelled": cancelled, "header": header, "rows": rows}


def export_datasets(datasets, output_folder, settings=None, prefix=None,
                    progress_callback=None, cancel_callback=None):
    """
    Features for several datasets in one file. datasets is a list of
    (pairs, sensor_file); each dataset is correlated against its own sensor file,
    and all of them go into one <prefix>_roi_metrics.csv and .xlsx in output_folder.

    Rows are written to the CSV as they are computed; the sensor columns are added
    when every dataset is done, so a stopped run still leaves the features.
    Returns the same keys as export_pairs, plus "sensor_columns".
    """
    import pandas as pd

    settings = merged_settings(settings)
    extractor = RoiFeatureExtractor(settings)
    header = extractor.header()
    prefix = prefix or run_prefix()
    os.makedirs(output_folder, exist_ok=True)
    csv_path = os.path.join(output_folder, f"{prefix}_roi_metrics.csv")

    total = sum(len(pairs) for pairs, _ in datasets)
    done, n_ok, n_failed, cancelled = 0, 0, 0, False
    frames = []

    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for pairs, sensor_file in datasets:
            rows = []
            for image_path, mask_path in pairs:
                if cancel_callback is not None and cancel_callback():
                    cancelled = True
                    break
                row = extractor.compute_row(image_path, mask_path)
                writer.writerow(row)
                f.flush()
                rows.append(row)
                n_ok, n_failed = (n_ok + 1, n_failed) if row[-1] == STATUS_OK else (n_ok, n_failed + 1)
                done += 1
                if progress_callback is not None:
                    progress_callback(done, total)
            if rows:
                frames.append(_with_sensor_columns(pd.DataFrame(rows, columns=header),
                                                   sensor_file, settings))
            if cancelled:
                break

    if not frames:
        return {"csv_path": csv_path, "xlsx_path": None, "rows_ok": n_ok, "rows_failed": n_failed,
                "cancelled": cancelled, "header": header, "rows": [], "sensor_columns": []}

    df = pd.concat(frames, ignore_index=True, sort=False)
    sensor_columns = [c for c in df.columns if c not in header]
    header = list(df.columns)
    rows = df.astype(object).where(df.notna(), "").values.tolist()
    df.to_csv(csv_path, index=False)

    xlsx_path = None
    if settings["output"]["excel"]:
        xlsx_path = os.path.join(output_folder, f"{prefix}_roi_metrics.xlsx")
        try:
            write_xlsx(xlsx_path, header, rows)
        except Exception as e:
            print(f"[roi_feature_extraction] Could not write {xlsx_path}: {e}")
            xlsx_path = None

    return {"csv_path": csv_path, "xlsx_path": xlsx_path, "rows_ok": n_ok, "rows_failed": n_failed,
            "cancelled": cancelled, "header": header, "rows": rows, "dataframe": df,
            "sensor_columns": sensor_columns}


def _with_sensor_columns(df, sensor_file, settings):
    """Append the nearest sensor reading for each image. Returns df unchanged on failure."""
    if not sensor_file or not os.path.isfile(sensor_file):
        return df
    try:
        import pandas as pd
        from appcore.SensorImageCorrelator import SensorImageCorrelator
        tolerance = settings["sensor"]["tolerance_minutes"] or None      # 0: automatic
        sensor_df = SensorImageCorrelator().sensor_values_for_images(
            df["Image Path"].tolist(), sensor_file, tolerance_minutes=tolerance)
        return pd.concat([df.reset_index(drop=True), sensor_df.reset_index(drop=True)], axis=1)
    except Exception as e:
        print(f"[roi_feature_extraction] Sensor correlation failed for {sensor_file}: {e}")
        return df


def segmentation_pairs(input_dir, output_dir, image_filter=None, include_missing=True):
    """
    (image, mask) pairs after segmenting input_dir, where masks are saved as
    <image stem>_mask.png in output_dir (a folder, or a list of folders searched in
    order). With include_missing, images without a mask are listed with mask None so
    the export records them as failed rows.
    """
    output_dirs = [output_dir] if isinstance(output_dir, str) else list(output_dir)
    pairs = []
    for name in sorted(os.listdir(input_dir)):
        stem, ext = os.path.splitext(name)
        if ext.lower() not in IMAGE_EXTENSIONS:
            continue
        if stem.lower().endswith(("_mask", "_overlay")):
            continue
        if image_filter is not None and name not in image_filter:
            continue
        mask = next((p for p in (os.path.join(d, f"{stem}_mask.png") for d in output_dirs)
                     if os.path.exists(p)), None)
        if mask is not None:
            pairs.append((os.path.join(input_dir, name), mask))
        elif include_missing:
            pairs.append((os.path.join(input_dir, name), None))
    return pairs
