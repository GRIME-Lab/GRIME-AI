#!/usr/bin/env python3
# -*- coding: utf-8 -*-

# Author: John Edward Stranzl, Jr.
# Affiliation(s): University of Nebraska-Lincoln, Blade Vision Systems, LLC
# Contact: jstranzl2@huskers.unl.edu, johnstranzl@gmail.com
# Created: Mar 6, 2022
# License: Apache License, Version 2.0, http://www.apache.org/licenses/LICENSE-2.0

import os
import datetime
import cv2
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from appcore.QProgressWheel import QProgressWheel
from appcore.App_Utils import App_Utils
from appcore.TimeStamp_Utils import TimeStamp_Utils
from appcore.Color import Color
from appcore.vegetation_indices import Vegetation_Indices
from appcore.dialogs.color_segmentation.color_seg_roi_data import roiData, roi_patch_and_mask

from appcore.Texture import GLCMTexture, LBPTexture, GaborTexture, WaveletTexture, FourierTexture

# ======================================================================================================================
# Plain-data ROI description. Qt objects stay in the main process; workers get
# only name, cluster count, rectangle and polygon as Python primitives.
# ======================================================================================================================
def _roi_specs(roiList):
    specs = []
    for roiObj in roiList:
        r = roiObj.getImageROI()
        shape = roiObj.getROIShape()
        is_rect = getattr(shape, 'value', shape) == 0
        poly = [] if is_rect else (roiObj.getImagePolygon() or [])
        specs.append({
            'name': roiObj.getROIName(),
            'nClusters': roiObj.getNumColorClusters(),
            'rect': (int(r.x()), int(r.y()), int(r.width()), int(r.height())),
            'shape': getattr(shape, 'value', shape),
            'polygon': [(int(p.x()), int(p.y())) for p in poly],
        })
    return specs


def load_image_mask(image_path, mask_spec, shape):
    """Per-image mask: <image stem>_mask.<ext> in mask_spec['folder'].

    mask_spec: {'name': column prefix, 'folder': mask folder, 'value': None or class index}
      value None -> binary mask: every non-zero pixel is inside the region
      value N    -> class-index mask: only pixels equal to N are inside

    Returns (mask, reason). mask is a boolean array the same size as the image, or
    None with a short reason when the mask is missing, the wrong size, or empty.
    """
    import os as _os
    import cv2 as _cv2
    stem = _os.path.splitext(_os.path.basename(image_path))[0]
    path = None
    for ext in ('.png', '.PNG', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp'):
        candidate = _os.path.join(mask_spec['folder'], stem + '_mask' + ext)
        if _os.path.isfile(candidate):
            path = candidate
            break
    if path is None:
        return None, f"no mask file '{stem}_mask.*'"

    m = _cv2.imread(path, _cv2.IMREAD_UNCHANGED)
    if m is None:
        return None, f"could not read {_os.path.basename(path)}"
    if m.ndim == 3:
        m = _cv2.cvtColor(m[:, :, :3], _cv2.COLOR_BGR2GRAY)
    if m.shape[:2] != tuple(shape[:2]):
        return None, (f"mask is {m.shape[1]}x{m.shape[0]}, image is {shape[1]}x{shape[0]}")

    value = mask_spec.get('value')
    mask = (m == value) if value is not None else (m != 0)
    if not mask.any():
        return None, 'mask is empty'
    return mask, None


def _crop_roi(img, spec):
    """(patch, mask): the ROI bounding box and a boolean mask of the pixels inside the ROI.
    (None, None) if the ROI does not overlap the image."""
    return roi_patch_and_mask(img, spec['rect'], spec['polygon'])


# ======================================================================================================================
# Masked texture. Every value is computed only from pixels inside the ROI mask:
#   GLCM            - only pixel pairs with both pixels inside
#   Gabor, LBP      - only responses whose whole filter footprint is inside
#   Wavelet (Haar)  - only coefficients whose whole support block is inside
#   Fourier         - largest axis-aligned rectangle fully inside (the transform is global)
# A method that has too few valid samples returns None (written as -999).
# ======================================================================================================================
_MIN_TEXTURE_SAMPLES = 16
_MIN_FOURIER_SIDE = 8


def _erode(mask, kh, kw):
    import cv2 as _cv2
    import numpy as _np
    k = _np.ones((kh, kw), _np.uint8)
    return _cv2.erode(mask.astype(_np.uint8), k, borderType=_cv2.BORDER_CONSTANT, borderValue=0).astype(bool)


def _masked_glcm_contrast(g, m):
    import numpy as _np
    from appcore.Texture import GLCMTexture
    t = GLCMTexture()
    H, W = g.shape
    vals = []
    for d in t.distances:
        for a in t.angles:
            dr = int(round(_np.sin(a) * d))
            dc = int(round(_np.cos(a) * d))
            r0, r1 = max(0, -dr), min(H, H - dr)
            c0, c1 = max(0, -dc), min(W, W - dc)
            if r1 <= r0 or c1 <= c0:
                continue
            A = g[r0:r1, c0:c1]
            B = g[r0 + dr:r1 + dr, c0 + dc:c1 + dc]
            v = m[r0:r1, c0:c1] & m[r0 + dr:r1 + dr, c0 + dc:c1 + dc]
            if v.sum() < _MIN_TEXTURE_SAMPLES:
                return None
            diff = A[v].astype(_np.float64) - B[v].astype(_np.float64)
            vals.append(float(_np.mean(diff * diff)))
    return float(_np.mean(vals)) if vals else None


def _masked_gabor_rms(g, m):
    """RMS of the Gabor response magnitude, sqrt(mean(real^2 + imag^2)), over valid pixels,
    for each frequency/orientation; returns the mean of those RMS values."""
    import numpy as _np
    from skimage.filters import gabor, gabor_kernel
    from appcore.Texture import GaborTexture
    t = GaborTexture()
    gf = g.astype(_np.float64)   # integer input would make skimage return a truncated integer response
    rms = []
    for f in t.frequencies:
        for th in t.thetas:
            kh, kw = gabor_kernel(f, theta=th).shape
            valid = _erode(m, kh, kw)
            if valid.sum() < _MIN_TEXTURE_SAMPLES:
                return None
            real, imag = gabor(gf, frequency=f, theta=th)
            rms.append(float(_np.sqrt(_np.mean(real[valid] ** 2 + imag[valid] ** 2))))
    return float(_np.mean(rms)) if rms else None


def _masked_lbp_entropy(g, m):
    import numpy as _np
    from skimage.feature import local_binary_pattern
    from appcore.Texture import LBPTexture
    t = LBPTexture()
    lbp = local_binary_pattern(g, t.P, t.R, method=t.method)
    k = 2 * int(_np.ceil(t.R)) + 1
    codes = lbp[_erode(m, k, k)]
    if codes.size < _MIN_TEXTURE_SAMPLES:
        return None
    p = _np.bincount(codes.astype(_np.int64)) / float(codes.size)
    p = p[p > 0]
    return float(-_np.sum(p * _np.log2(p))) + 0.0   # + 0.0 turns -0.0 into 0.0


def _block_all(m, block):
    """True where every pixel of a block x block cell is inside (cells past the edge count as outside)."""
    import numpy as _np
    H, W = m.shape
    Hb, Wb = -(-H // block), -(-W // block)
    pad = _np.zeros((Hb * block, Wb * block), dtype=bool)
    pad[:H, :W] = m
    return pad.reshape(Hb, block, Wb, block).all(axis=(1, 3))


def _masked_wavelet_detail_var(g, m):
    import numpy as _np
    import pywt
    from appcore.Texture import WaveletTexture
    t = WaveletTexture()
    if pywt.Wavelet(t.wavelet).dec_len != 2:
        print(f'[texture] masked wavelet supports Haar-length wavelets only (got {t.wavelet}).')
        return None
    coeffs = pywt.wavedec2(g.astype(_np.float32), wavelet=t.wavelet, level=t.level)
    vars_ = []
    # coeffs[1] is the coarsest detail level (t.level), coeffs[-1] the finest (1)
    for k, (cH, cV, cD) in enumerate(coeffs[1:]):
        lvl = t.level - k
        valid = _block_all(m, 2 ** lvl)
        h = min(valid.shape[0], cH.shape[0])
        w = min(valid.shape[1], cH.shape[1])
        v = valid[:h, :w]
        if v.sum() < _MIN_TEXTURE_SAMPLES:
            return None
        for c in (cH, cV, cD):
            vars_.append(float(_np.var(c[:h, :w][v])))
    return float(_np.mean(vars_)) if vars_ else None


def _largest_inside_rect(m):
    """(r0, r1, c0, c1) of the largest axis-aligned rectangle of True pixels."""
    import numpy as _np
    H, W = m.shape
    if m.all():
        return 0, H, 0, W
    heights = _np.zeros(W, dtype=_np.int64)
    best = (0, 0, 0, 0, 0)   # area, r0, r1, c0, c1
    for r in range(H):
        heights = _np.where(m[r], heights + 1, 0)
        stack = []
        for c in range(W + 1):
            h = heights[c] if c < W else 0
            start = c
            while stack and stack[-1][1] >= h:
                sc, sh = stack.pop()
                area = sh * (c - sc)
                if area > best[0]:
                    best = (area, r - sh + 1, r + 1, sc, c)
                start = sc
            stack.append((start, h))
    return best[1], best[2], best[3], best[4]


def _masked_fourier_mean(g, m):
    import numpy as _np
    from appcore.Texture import FourierTexture
    r0, r1, c0, c1 = _largest_inside_rect(m)
    if (r1 - r0) < _MIN_FOURIER_SIDE or (c1 - c0) < _MIN_FOURIER_SIDE:
        return None
    prof = _np.asarray(FourierTexture().compute_features(g[r0:r1, c0:c1]), dtype=float)
    return float(_np.mean(prof)) if prof.size else None


# ======================================================================================================================
# Column layout. One layout drives the CSV header, every CSV row and every XLSX sheet, so they
# always line up. A layout item is (group, sheet, column, key):
#   group  - region the column belongs to ('whole' or 'roi<i>'); a blank column separates groups
#   sheet  - XLSX sheet
#   column - column title
#   key    - lookup key in that region's computed values
# ======================================================================================================================
_TEXTURE_METHODS = (('glcm', 'GLCM_contrast'), ('gabor', 'Gabor_RMS'), ('lbp', 'LBP_entropy'),
                    ('wavelet', 'Wavelet_detail_var'), ('fourier', 'Fourier_radial_mean'))
SHEET_WHOLE = 'Whole Image'
SHEET_FEATURES = 'ROI Features'
SHEET_TEXTURE = 'ROI Texture'
SHEET_GREENNESS = 'ROI Greenness'
SHEET_COLORS = 'ROI Colors'
XLSX_SHEETS = (SHEET_WHOLE, SHEET_FEATURES, SHEET_TEXTURE, SHEET_GREENNESS, SHEET_COLORS)
WHOLE_LABEL = 'Whole Image'
# Fill for the blank separator columns in the XLSX (light gray). CSV has no colors.
SEPARATOR_FILL = 'E7E6E6'
_PLACEHOLDER = -999


def _flags_from_params(p):
    return {
        'wholeImage': bool(p.wholeImage),
        'ROI': bool(p.ROI),
        'Intensity': bool(p.Intensity),
        'ShannonEntropy': bool(p.ShannonEntropy),
        'Texture': bool(p.Texture),
        'HSV': bool(p.HSV),
    }


def _texture_labels(flags, texture_options):
    if not flags.get('Texture'):
        return []
    opts = texture_options or {}
    return [label for key, label in _TEXTURE_METHODS if opts.get(key)]


def _region_layout(group, label, is_whole, nClusters, flags, greenness_names, texture_options,
                   is_mask=False):
    """Columns for one region in CSV order: Intensity, Entropy, texture, greenness, colors.
    Whole image: everything on the Whole Image sheet with plain names.
    ROI: ROI Features / ROI Texture / ROI Greenness / ROI Colors sheets, prefixed with the ROI name.
    Colors are grouped per cluster (H, S, V, Coverage)."""
    p = '' if is_whole else f'{label}: '
    s_feat = SHEET_WHOLE if is_whole else SHEET_FEATURES
    s_tex = SHEET_WHOLE if is_whole else SHEET_TEXTURE
    s_grn = SHEET_WHOLE if is_whole else SHEET_GREENNESS
    s_col = SHEET_WHOLE if is_whole else SHEET_COLORS
    items = []
    if is_mask:
        # The masked area changes from image to image, so it is a feature in its own right.
        items.append((group, s_feat, p + 'Area_px', 'Area_px'))
        items.append((group, s_feat, p + 'Area_frac', 'Area_frac'))
    if flags.get('Intensity'):
        items.append((group, s_feat, p + 'Intensity', 'Intensity'))
    if flags.get('ShannonEntropy'):
        items.append((group, s_feat, p + 'Entropy', 'Entropy'))
    for t in _texture_labels(flags, texture_options):
        items.append((group, s_tex, p + t, 'tex:' + t))
    for g in greenness_names:
        items.append((group, s_grn, p + g, 'g:' + g))
    if flags.get('HSV'):
        for i in range(nClusters):
            for k in ('H', 'S', 'V', 'Coverage'):
                items.append((group, s_col, f'{p}{k}_{i}', f'{k}_{i}'))
    return items


def _layout(nClusters, flags, greenness_names, texture_options, roi_specs, mask_spec=None):
    items = []
    if flags.get('wholeImage'):
        items += _region_layout('whole', WHOLE_LABEL, True, nClusters, flags, greenness_names, texture_options)
    if flags.get('ROI'):
        for i, spec in enumerate(roi_specs):
            items += _region_layout(f'roi{i}', spec['name'], False, spec['nClusters'], flags,
                                    greenness_names, texture_options)
    if mask_spec:
        items += _region_layout('mask', mask_spec['name'], False, nClusters, flags,
                                greenness_names, texture_options, is_mask=True)
    return items


def _fmt(key, value):
    if isinstance(value, int) and value == _PLACEHOLDER:
        return '%d' % value
    if key[:2] in ('H_', 'S_', 'V_'):
        return '%3.2f' % value
    return '%3.4f' % value


def _csv_line(first, fields):
    """first: leading text (hyperlink,date,time or header titles). fields: [(group, text)].
    A blank column is inserted wherever the group changes."""
    parts = [first]
    prev = None
    for group, text in fields:
        if prev is not None and group != prev:
            parts.append(',')
        parts.append(', ' + text)
        prev = group
    return ''.join(parts) + '\n'


def _region_values(helper, color, rgb, gray, gray_tex, mask, nClusters, flags, greenness_list, texture_options):
    """Feature values for one region, keyed like _region_layout.
    rgb/gray: the region's pixels (whole image, or an (N, 1, C) strip of inside pixels).
    gray_tex/mask: 2-D grayscale for texture and its inside-mask (None = whole patch)."""
    import cv2 as _cv2
    from appcore.vegetation_indices import Vegetation_Indices
    vals = {}
    if flags.get('Intensity'):
        # The range for a pixel's value in grayscale is (0-255), 127 lies midway
        vals['Intensity'] = float(_cv2.mean(gray)[0])
    if flags.get('ShannonEntropy'):
        vals['Entropy'] = helper.calcEntropy(gray)
    if flags.get('Texture'):
        for n, v in helper.compute_texture_scalars(gray_tex, texture_options, mask):
            vals['tex:' + n] = float(v)
    for g in greenness_list:
        vals['g:' + g.get_name()] = float(Vegetation_Indices().get_greenness(g, rgb).get_value())
    if flags.get('HSV'):
        # Dominant colors, largest coverage first. Coverage = fraction of the region's pixels (0..1).
        hist, centers = color.extractDominant_HSV(rgb, nClusters)
        centers = np.asarray(centers, dtype=np.float32).reshape(-1, 3).copy()
        # OpenCV stores hue as 0..180 to fit a byte; report the true 0..360 color-space value.
        centers[:, 0] *= 2.0
        hist = [] if hist is None else list(np.asarray(hist, dtype=float).ravel())
        for i in range(nClusters):
            if i < len(centers) and i < len(hist):
                vals[f'H_{i}'], vals[f'S_{i}'], vals[f'V_{i}'] = (float(x) for x in centers[i][:3])
                vals[f'Coverage_{i}'] = float(hist[i])
            else:
                # k-means returned fewer distinct clusters than requested (e.g. a uniform region)
                vals[f'H_{i}'] = vals[f'S_{i}'] = vals[f'V_{i}'] = _PLACEHOLDER
                vals[f'Coverage_{i}'] = 0.0
    return vals


# ======================================================================================================================
# Build one image's record.
# ======================================================================================================================
def _build_feature_row(helper, color, path, nClusters, flags, greenness_list, texture_options, roi_specs,
                       mask_spec=None):
    """Returns (csv_row, record). record = {'path', 'date', 'time', 'cells': [(group, sheet, column, value)]}."""
    import cv2 as _cv2
    from appcore.TimeStamp_Utils import TimeStamp_Utils

    link = helper.create_hyperlink(path)
    ts = TimeStamp_Utils(); ts.detectDateTime(path)
    d, t = ts.extractDateTime(path)

    img = color.loadColorImage(path)   # RGB
    values = {}

    if flags.get('wholeImage'):
        gray = _cv2.cvtColor(img, _cv2.COLOR_RGB2GRAY)
        values['whole'] = _region_values(helper, color, img, gray, gray, None, nClusters, flags,
                                         greenness_list, texture_options)

    if flags.get('ROI') and roi_specs:
        values.update(helper.calculate_ROI_scalars(img, roi_specs, flags, greenness_list, texture_options))

    if mask_spec:
        # One mask per image: features are measured only inside it.
        mask, reason = load_image_mask(path, mask_spec, img.shape)
        if mask is None:
            print(f"[feature-export] {os.path.basename(path)}: {reason}; writing -999.")
        else:
            rgb = img[mask].reshape(-1, 1, 3)
            gray = _cv2.cvtColor(rgb, _cv2.COLOR_RGB2GRAY)
            gray_tex = _cv2.cvtColor(img, _cv2.COLOR_RGB2GRAY) if flags.get('Texture') else None
            n_inside = int(mask.sum())
            if n_inside < (4 * max(1, nClusters) if flags.get('HSV') else 1):
                print(f"[feature-export] {os.path.basename(path)}: mask has {n_inside} pixel(s); writing -999.")
            else:
                vals = _region_values(helper, color, rgb, gray, gray_tex, mask, nClusters, flags,
                                      greenness_list, texture_options)
                vals['Area_px'] = float(n_inside)
                vals['Area_frac'] = float(n_inside) / float(img.shape[0] * img.shape[1])
                values['mask'] = vals

    layout = _layout(nClusters, flags, [g.get_name() for g in greenness_list], texture_options, roi_specs, mask_spec)
    cells = []
    fields = []
    for group, sheet, col, key in layout:
        v = values.get(group, {}).get(key, _PLACEHOLDER)
        cells.append((group, sheet, col, v))
        fields.append((group, _fmt(key, v)))

    row = _csv_line('%s,%s,%s' % (link, d, t), fields)
    record = {'path': path, 'date': d, 'time': t, 'cells': cells}
    return row, record


# ======================================================================================================================
# Module-level worker. Must be top-level (picklable) so ProcessPoolExecutor can
# dispatch it. Returns (index, csv_row, record) so rows are written in original order.
# Also used by the serial fallback so both paths produce identical rows.
# ======================================================================================================================
def _compute_image_row(args):
    (index, path, nClusters, flags, greenness_list, texture_options, roi_specs, mask_spec) = args
    import os as _os
    from appcore.Color import Color
    from appcore.dialogs.color_segmentation.color_seg_feature_export import ColorSegFeatureExport
    if not _os.path.isfile(path):
        return (index, None, None)
    try:
        row, record = _build_feature_row(ColorSegFeatureExport(), Color(), path, nClusters,
                                         flags, greenness_list, texture_options, roi_specs, mask_spec)
        return (index, row, record)
    except Exception as e:
        print(f'[feature-export worker] {path}: {e}')
        return (index, None, None)


# ======================================================================================================================
# XLSX: Whole Image, ROI Features, ROI Texture, ROI Greenness and ROI Colors sheets (sheets with no columns are
# omitted). Every sheet starts with Image (hyperlinked), Date and Time; one row per image.
# A blank column separates regions (whole image, each ROI).
# ======================================================================================================================
def _write_feature_xlsx(xlsx_path, records):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    records = [r for r in records if r]
    if not records:
        return None

    # Column plan per sheet from the (identical) layout of every record: names, with None = blank separator.
    plan = {sh: [] for sh in XLSX_SHEETS}
    last_group = {sh: None for sh in XLSX_SHEETS}
    for group, sheet, col, _ in records[0]['cells']:
        if last_group[sheet] is not None and group != last_group[sheet]:
            plan[sheet].append(None)
        plan[sheet].append(col)
        last_group[sheet] = group

    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    link_font = Font(color='0563C1', underline='single')
    sep_fill = PatternFill(fill_type='solid', start_color=SEPARATOR_FILL, end_color=SEPARATOR_FILL)
    for sheet in XLSX_SHEETS:
        cols = plan[sheet]
        if not cols:
            continue
        ws = wb.create_sheet(sheet)
        ws.freeze_panes = 'D2'
        ws.column_dimensions['A'].width = 45
        ws.column_dimensions['B'].width = 12
        ws.column_dimensions['C'].width = 10
        for j, name in enumerate(cols, start=4):
            ws.column_dimensions[get_column_letter(j)].width = 3 if name is None else max(12, min(40, len(name) + 2))

        def sep_cell():
            c = WriteOnlyCell(ws, value=None)
            c.fill = sep_fill
            return c

        header = []
        for name in ['Image', 'Date (ISO)', 'Time (ISO)'] + cols:
            if name is None:
                header.append(sep_cell())
                continue
            c = WriteOnlyCell(ws, value=name)
            c.font = bold
            header.append(c)
        ws.append(header)

        for r in records:
            img_cell = WriteOnlyCell(ws, value=os.path.basename(r['path']))
            img_cell.hyperlink = r['path']
            img_cell.font = link_font
            row = [img_cell, r['date'], r['time']]
            prev = None
            for group, sh, _, v in r['cells']:
                if sh != sheet:
                    continue
                if prev is not None and group != prev:
                    row.append(sep_cell())
                row.append(v)
                prev = group
            ws.append(row)

    wb.save(xlsx_path)
    return xlsx_path


class ColorSegFeatureExport:
    def __init__(self):
        csvFilename = ''
        self.instance = 1

    # ==================================================================================================================
    #
    # ==================================================================================================================
    def create_training_data_filename(self, imageFileFolder):
        base = 'TrainingData_' + datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        imageQualityFile = os.path.join(imageFileFolder, base + '.csv')
        csvFilename = open(imageQualityFile, 'a', newline='')

        # Same base name as the CSV so the pair is easy to find.
        xlsxFilename = os.path.join(imageFileFolder, base + '.xlsx')

        return csvFilename, xlsxFilename

    # ==================================================================================================================
    #
    # ==================================================================================================================
    def ExtractFeatures(self, imagesList, imageFileFolder, roiList, colorSegmentationParams, greenness_index_list, texture_options=None, mask_spec=None):

        # Region Select: Whole Image and/or ROI (fixed ROIs or a per-image mask).
        # Nothing to extract if none applies.
        do_whole = bool(colorSegmentationParams.wholeImage)
        do_roi = bool(colorSegmentationParams.ROI) and (len(roiList) > 0 or bool(mask_spec))
        if not (do_whole or do_roi):
            print('[feature-export] Nothing to extract: select Whole Image, '
                  'or select ROI and add at least one ROI.')
            return

        # ----------------------------------------------------------------------------------------------------------
        # CREATE PROGRESS WHEEL
        # ----------------------------------------------------------------------------------------------------------
        myGRIMe_Color = Color()

        # GENERATE LIST OF IMAGE FILES IN FOLDER
        videoFileList = imagesList
        nFrameCount = len(videoFileList)

        progressBar = QProgressWheel(0, len(videoFileList) + 1)
        progressBar.show()
        progressBarIndex = 1

        if nFrameCount > 0:
            # CREATE TRAINING DATA FILENAME
            csvFile, xlsFile = self.create_training_data_filename(imageFileFolder)

            # CREATE THE OUTPUT FILE COLUMN HEADER
            nMaxNumColorClusters = App_Utils().getMaxNumColorClusters(roiList)

            '''
            for roiObj in roiList:

                # Build the initial columns
                header_cols = ["Training Filename", "Date (ISO)", "Time (ISO)",  "Intensity", "Entropy"]
                base_header = ",".join(header_cols)

                for greenness in greenness_index_list:
                    base_header = base_header + ',' + greenness.get_name()

                base_header = base_header + ',,'

                # Create the template for HSV headings based on the ROI name
                template = f", {roiObj.getROIName()}: #"

                # Build the HSV part of the header
                if nMaxNumColorClusters == 1:
                    hsv_cols = ''.join(template.replace('#', char) for char in ['H', 'S', 'V'])
                else:
                    hsv_cols = ''.join(
                        ''.join(f"{template.replace('#', char)}_{idx}" for idx in range(nMaxNumColorClusters))
                        for char in ['H', 'S', 'V']
                    )

                # Combine all parts and add a newline at the end
                strOutputString = base_header + hsv_cols + "\n"
                strOutputString += "\n"

                csvFile.write(strOutputString)

                # LOAD THE IMAGE FILE
                img = myGRIMe_Color.loadColorImage(roiObj.getTrainingImageName())

                # CONVERT TO GRAY SCALE
                gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
                height, width = gray.shape
                gray = cv2.resize(gray, (width, height))

                texture = -999

                # CREATE HYPERLINK TO FILE
                strOutputString = self.create_hyperlink(roiObj.getTrainingImageName())
                strOutputString = strOutputString + ','     # DATE
                strOutputString = strOutputString + ','     # TIME

                # COMPUTE INTENSITY FOR THE ENTIRE IMAGE
                # ------------------------------------------------------------------------------------------------------
                intensity = cv2.mean(gray)[0]  # The range for a pixel's value in grayscale is (0-255), 127 lies midway
                strOutputString = strOutputString + ',' + '%3.4f' % intensity

                # COMPUTE ENTROPY FOR THE ENTIRE IMAGE
                # ------------------------------------------------------------------------------------------------------
                entropyValue = self.calcEntropy(gray)
                strOutputString = strOutputString + ',' + '%3.4f' % entropyValue

                # CALCULATE THE GREENNESS INDEX FOR THE ROI
                # ------------------------------------------------------------------------------------------------------
                try:
                    for greenness in greenness_index_list:
                        greenness_updated = Vegetation_Indices().get_greenness(greenness, img)
                        strOutputString = strOutputString + ',' + '%3.4f' % greenness_updated.get_value()
                except:
                    pass

                # SKIP TWO COLUMNS TO ALIGN THE WHOLE IMAGE HEADER WITH THE ROI HEADERS
                strOutputString = strOutputString + ',,'

                hsvClusterCenters, hist = roiObj.getHSVClusterCenters()
                for i in range(nMaxNumColorClusters):
                    strOutputString = strOutputString + ',' + str(hsvClusterCenters[i][0]) + ',' + str(hsvClusterCenters[i][1]) + ',' + str(hsvClusterCenters[i][2])
                strOutputString = strOutputString + '\n'

                csvFile.write(strOutputString)
            '''


            self.extract_ROI_features(csvFile, xlsFile, videoFileList, roiList, colorSegmentationParams, greenness_index_list, progressBarIndex, progressBar, texture_options=texture_options, mask_spec=mask_spec)


            csvFile.close()

            # CLOSE AND DELETE THE PROGRESSBAR
            progressBar.close()
            del progressBar

    # ==================================================================================================================
    #
    # ==================================================================================================================
    def build_scalar_header(self, nClusters, roiList, colorSegmentationParams, greenness_index_list, texture_options=None, mask_spec=None):
        """CSV header, built from the same layout as the data rows."""
        flags = _flags_from_params(colorSegmentationParams)
        specs = _roi_specs(roiList) if flags['ROI'] else []
        layout = _layout(nClusters, flags, [g.get_name() for g in greenness_index_list], texture_options, specs,
                         mask_spec if flags['ROI'] else None)
        return _csv_line('Image, Date (ISO), Time (ISO)', [(grp, col) for grp, _, col, _ in layout])


    # ==================================================================================================================
    #
    # ==================================================================================================================
    def extract_ROI_features(self, csvFile, xlsFile, videoFileList, roiList, colorSegmentationParams, greenness_index_list, progressBarIndex, progressBar, texture_options=None, mask_spec=None):

        myGRIMe_Color = Color()

        # ----------------------------------------------------------------------------------------------------------
        # PROCESS IMAGES
        # ----------------------------------------------------------------------------------------------------------
        # E:\\000 - University of Nebraska\\2022_USGS104b\\imagery\\PBT_MittelstetsMeanderCam\\20220709

        # ONCE THE NUMBER OF COLOR CLUSTERS HAS BEEN SELECTED AND AN ROI TRAINED, LOCK THE NUMBER OF COLOR CLUSTERS
        # AND TRAIN ALL SUBSEQUENT ROIs FOR THE SAME NUMBER OF COLOR CLUSTERS
        nClusters = colorSegmentationParams.numColorClusters

        texture_options = texture_options or {}
        header = self.build_scalar_header(nClusters, roiList, colorSegmentationParams, greenness_index_list, texture_options, mask_spec)

        # WRITE THE HEADER TO THE CSV
        csvFile.write(header)

        flags = _flags_from_params(colorSegmentationParams)
        roi_specs = _roi_specs(roiList) if flags['ROI'] else []
        mask_spec = mask_spec if flags['ROI'] else None
        greenness_list = list(greenness_index_list)
        tasks = [(i, fname.fullPathAndFilename, nClusters, flags, greenness_list, texture_options, roi_specs, mask_spec)
                 for i, fname in enumerate(videoFileList)]

        # ------------------------------------------------------------------
        # PARALLEL: every image's row is independent and CPU-bound, so compute
        # across a process pool (whole image, ROIs, or both). Rows are written
        # in original order.
        # ------------------------------------------------------------------
        from concurrent.futures import ProcessPoolExecutor
        import os as _os
        results = [None] * len(tasks)
        workers = max(1, (_os.cpu_count() or 2) - 1)
        chunk = max(1, len(tasks) // (workers * 4))
        print(f'[feature-export] PARALLEL: {workers} workers, {len(tasks)} images, '
              f'whole image {"on" if flags["wholeImage"] else "off"}, {len(roi_specs)} ROIs'
              + (f", per-image mask '{mask_spec['name']}'" if mask_spec else ''))
        done = 0
        parallel_ok = False
        try:
            with ProcessPoolExecutor(max_workers=workers) as ex:
                for index, row, record in ex.map(_compute_image_row, tasks, chunksize=chunk):
                    results[index] = (row, record)
                    done += 1
                    progressBar.setValue(progressBarIndex + done)
            parallel_ok = True
        except Exception as _e:
            print(f'[feature-export] parallel path failed ({_e}); falling back to serial.')

        if not parallel_ok:
            # ------------------------------------------------------------------
            # SERIAL FALLBACK: same row builder, in-process.
            # ------------------------------------------------------------------
            print(f'[feature-export] SERIAL: {len(tasks)} images')
            results = [None] * len(tasks)
            for task in tasks:
                index, row, record = _compute_image_row(task)
                results[index] = (row, record)
                progressBarIndex += 1
                progressBar.setValue(progressBarIndex)

        for row, _ in results:
            if row:
                csvFile.write(row)

        try:
            out = _write_feature_xlsx(xlsFile, [rec for _, rec in results])
            if out:
                print(f'[feature-export] XLSX written: {out}')
        except Exception as _e:
            print(f'[feature-export] XLSX not written ({_e}).')



    # ==================================================================================================================
    #
    # ==================================================================================================================
    # ==================================================================================================================
    def compute_texture_scalars(self, gray, texture_options, mask=None):
        """Compute one representative scalar per SELECTED texture method on a
        grayscale image. Returns an ordered list of (name, value) so the CSV
        header and the data row always match. Only methods whose checkbox is
        on are included. If mask is given, only pixels where mask is True are
        used (see the masked helpers above); -999 if a method has too few
        valid samples."""
        import numpy as _np
        import cv2 as _cv2b
        opts = texture_options or {}
        g = gray
        m = _np.ones(g.shape[:2], dtype=bool) if mask is None else mask.astype(bool)

        # Texture scalars are aggregate stats; full resolution is wasteful.
        # Downsample the longest side to <=512 px. A downsampled pixel stays
        # in the mask only if every source pixel it averages was inside.
        h, w = g.shape[:2]
        mx = max(h, w)
        if mx > 512:
            sc = 512.0 / mx
            size = (max(1, int(w * sc)), max(1, int(h * sc)))
            g = _cv2b.resize(g, size, interpolation=_cv2b.INTER_AREA)
            m = _cv2b.resize(m.astype(_np.float32), size, interpolation=_cv2b.INTER_AREA) >= 0.999

        funcs = {'glcm': _masked_glcm_contrast, 'gabor': _masked_gabor_rms, 'lbp': _masked_lbp_entropy,
                 'wavelet': _masked_wavelet_detail_var, 'fourier': _masked_fourier_mean}
        out = []
        for key, name in _TEXTURE_METHODS:
            fn = funcs[key]
            if not opts.get(key):
                continue
            try:
                v = fn(g, m)
            except Exception as _e:
                print(f'[texture] {name} error: {_e}')
                v = None
            out.append((name, -999.0 if v is None else v))
        return out

    # ==================================================================================================================
    # CALCULATE THE FEATURE VALUES FOR EACH ROI
    # ==================================================================================================================
    def calculate_ROI_scalars(self, img, roi_specs, flags, greenness_index_list, texture_options=None):
        """roi_specs: list of dicts from _roi_specs(). flags: dict of feature switches.
        Returns {'roi<i>': values} keyed like _region_layout. An ROI with too few pixels in this
        image gets no entry, so all its columns are written as -999."""
        out = {}
        texture_options = texture_options or {}
        color = Color()

        for i, spec in enumerate(roi_specs):
            nClusters = spec['nClusters']

            patch, mask = _crop_roi(img, spec)
            n_inside = 0 if mask is None else int(mask.sum())
            # color clustering subsamples every 4th pixel of an irregular region
            needed = 4 * max(1, nClusters) if flags.get('HSV') else 1
            if n_inside < needed:
                print(f"[feature-export] ROI '{spec['name']}' has {n_inside} pixels in this image; writing -999.")
                continue

            # Only pixels inside the ROI, as an (N, 1, 3) strip, for all non-texture features.
            rgb = patch[mask].reshape(-1, 1, 3)
            gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
            # Texture needs the 2-D patch; the mask excludes everything outside the ROI.
            gray_patch = cv2.cvtColor(patch, cv2.COLOR_RGB2GRAY) if flags.get('Texture') else None

            out[f'roi{i}'] = _region_values(self, color, rgb, gray, gray_patch, mask, nClusters, flags,
                                            greenness_index_list, texture_options)
        return out


    # ==================================================================================================================
    #
    # ==================================================================================================================
    def calcEntropy(self, img):
        entropy = []

        hist = cv2.calcHist([img], [0], None, [256], [0, 255])
        total_pixel = img.shape[0] * img.shape[1]

        for item in hist:
            probability = item / total_pixel
            if probability == 0:
                en = 0
            else:
                en = -1 * probability * (np.log(probability) / np.log(2))
            entropy.append(en)

        try:
            sum_en = sum(entropy)
        except Exception:
            sum_en = 0.0

        # hist items are 1-element arrays; return a plain float (NumPy 2 won't format arrays with %f).
        return float(np.asarray(sum_en).ravel()[0]) if np.ndim(sum_en) else float(sum_en)

    def create_hyperlink(self, file_path: str) -> str:
        """
        Returns an Excel HYPERLINK formula for the given file path.
        This function cleans the file path and escapes it appropriately for CSV.
        """
        # Clean the file path by stripping extra whitespace and any extraneous quotes
        clean_path = file_path.strip().strip('"')

        # Create the hyperlink formula
        formula = f'=HYPERLINK("{clean_path}", "{os.path.basename(clean_path)}")'

        # If the formula contains commas, enclose the entire field in extra quotes and escape inner quotes.
        if ',' in formula:
            safe_formula = formula.replace('"', '""')
            formula = f'"{safe_formula}"'

        return formula


    def extract_texture_and_save(self, image, excel_filename, plot_folder, file_name, plot_option="both"):
        """
        Process an image, extract texture features, save them to an Excel file, and handle Fourier profile plots.

        Parameters:
            image (numpy array): Input image data.
            excel_filename (str): Name of the Excel file to append texture features.
            plot_folder (str): Folder path to save Fourier profile plots (if required).
            file_name (str): Name of the image (used for labeling data in Excel).
            plot_option (str): Options for handling the plot:
                               - "file": Save the plot as an image file.
                               - "embed": Embed the plot into the Excel file.
                               - "both": Save the plot as an image file and embed it into the Excel file.
        """
        # Ensure the file exists
        if not os.path.exists(excel_filename):
            # Create the file with a default visible sheet
            with pd.ExcelWriter(excel_filename, engine="openpyxl") as writer:
                # Add a default empty DataFrame to create a visible sheet
                pd.DataFrame().to_excel(writer, sheet_name="Default_Sheet", index=False)

        # Ensure the output folder for plots exists if saving as files
        if plot_option in ["file", "both"]:
            os.makedirs(plot_folder, exist_ok=True)

        # Instantiate texture measurement objects
        glcm_obj = GLCMTexture()
        gabor_obj = GaborTexture()
        lbp_obj = LBPTexture()
        wavelet_obj = WaveletTexture()
        fourier_obj = FourierTexture()

        # Compute features using each method
        glcm_features = glcm_obj.compute_features(image)
        gabor_features = gabor_obj.compute_features(image)
        lbp_features = lbp_obj.compute_features(image)
        wavelet_features = wavelet_obj.compute_features(image)
        fourier_features = fourier_obj.compute_features(image)

        # Create DataFrames for the texture features
        glcm_df = pd.DataFrame({"GLCM Features": glcm_features})
        gabor_df = pd.DataFrame({"Gabor Features": gabor_features})
        lbp_df = pd.DataFrame({"LBP Histogram": lbp_features})
        wavelet_df = pd.DataFrame({"Wavelet Features": wavelet_features})
        fourier_df = pd.DataFrame({"Fourier Radial Profile": fourier_features})

        # Write DataFrames to the Excel file under separate sheets for the image
        try:
            with pd.ExcelWriter(excel_filename, mode="a", engine="openpyxl") as writer:
                glcm_df.to_excel(writer, sheet_name=f"{file_name}_GLCM")
                gabor_df.to_excel(writer, sheet_name=f"{file_name}_Gabor")
                lbp_df.to_excel(writer, sheet_name=f"{file_name}_LBP")
                wavelet_df.to_excel(writer, sheet_name=f"{file_name}_Wavelet")
                fourier_df.to_excel(writer, sheet_name=f"{file_name}_Fourier")
        except Exception:
            pass

        # Plot the F    self._book = load_workbook(self._handles.handle, **engine_kwargs)ourier radial profile
        plt.figure(figsize=(6, 4))
        plt.plot(fourier_features, marker='o')
        plt.title(f"Fourier Radial Profile - {file_name}")
        plt.xlabel("Radial Bin")
        plt.ylabel("Average Magnitude")
        plt.grid(True)

        # Handle plot based on user option
        plot_filename = os.path.join(plot_folder, f"{file_name}_fourier_profile.png")
        if plot_option in ["file", "both"]:
            plt.savefig(plot_filename)  # Save the plot to a file
        if plot_option in ["embed", "both"]:
            plt.savefig("temp_plot.png")  # Temporary file to embed the plot
            plt.close()

            # Open the workbook to embed the plot
            workbook = load_workbook(excel_filename)
            worksheet = workbook[f"{file_name}_Fourier"]
            img = ExcelImage("temp_plot.png")
            worksheet.add_image(img, "B10")  # Embed image at cell B10
            workbook.save(excel_filename)

        print(
            f"Texture features for {file_name} written to {excel_filename}. Plots handled as per option '{plot_option}'.")
